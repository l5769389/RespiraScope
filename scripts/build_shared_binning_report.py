"""Build the comparison DOCX from reviewed measurements and figures.
Run with the document workspace's bundled Python (python-docx).
"""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'docs/measurements/shared_binning'
OUT=ROOT/'docs/时间分箱与固定幅值分箱波形对比.docx'
rows=json.loads((DATA/'metrics.json').read_text())['rows']
# Standalone build: obsolete comparison documents are no longer dependencies.
d=Document()
sec=d.sections[0]
sec.page_width=Inches(11.69);sec.page_height=Inches(8.27)
sec.top_margin=sec.bottom_margin=Inches(.4);sec.left_margin=sec.right_margin=Inches(.4)
sec.header_distance=sec.footer_distance=Inches(.16)
for name,size in [('Normal',10.5),('Title',22),('Heading 1',14),('Caption',9)]:
    st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('000000' if name in ('Title','Heading 1') else '182230')
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Hiragino Sans GB')
    st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(5)
    st.paragraph_format.line_spacing=Pt(30 if name=='Title' else 19 if name=='Heading 1' else 12 if name=='Caption' else 15)
    st.font.bold=name=='Heading 1'
    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');st.element.get_or_add_pPr().append(snap)
    st.paragraph_format.keep_with_next=False
    for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.text='RespiraScope  ·  模拟测算  ·  '
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for r in footer.runs:r.font.size=Pt(8)

def para(t,style=None):return d.add_paragraph(t,style)

def table(headers,widths):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for c,tx in zip(t.rows[0].cells,headers):c.text=tx
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    return t

def style_table(t,widths,compact=False):
    for i,row in enumerate(t.rows):
        no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
        for c,w in zip(row.cells,widths):
            c.width=Inches(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'45' if compact and side in ['top','bottom'] else '75');e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar);borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF6');pr.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2 if compact else 4);p.paragraph_format.keep_with_next=False
                p.paragraph_format.line_spacing=1 if p._p.xpath('.//w:drawing | .//m:oMath') else Pt(12.5)
                for r in p.runs:
                    r.font.size=Pt(9.5)
                    if i==0:r.bold=True

def fmt(value):
    return '—' if value is None else f'{value:.2f}%'


def add_metric(cell,value):
    p=cell.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f'参考不一致率 {fmt(value)}')

def explain(label,text):
    p=para('', 'Caption');p.paragraph_format.space_after=Pt(3)
    p.add_run(label+'：').bold=True
    p.add_run(text)

widths=[.95,4.1,4.1,1.7]
for offset in range(0,len(rows),2):
    if offset:
        p=para(f'时间与固定幅值分箱波形对比 {offset+1} 至 {offset+2}','Heading 1')
        p.paragraph_format.page_break_before=True
    else:
        para('时间与固定幅值分箱波形对比','Title')
        para('16 种模拟波形，每种 20 次，每周期 10 箱；图示为其中一次。两种分箱共用滤波曲线。','Caption')
        explain('参考来源','用未加入噪声和伪影的模拟曲线，分别按时间规则、固定幅值规则生成参考分箱；幅值上下限与实际设置相同。')
        explain('核对数据占比','两种方式都已分箱且有参考划分的点数，占整段记录的比例。分母是整段记录的点数。')
        explain('参考不一致率','在这批共同核对点中，某方式的实际吸呼方向或箱号与自身参考不同的比例。分母是核对点数，不是统计两种方式彼此有多少不同。')
        explain('例','整段 1,000 点，共同核对 900 点，占比为 90%；其中时间分箱有 27 点与时间参考不同，不一致率为 3%（27 ÷ 900）。')
        para('灰线：原始；蓝线：滤波；绿线：无噪声参考；橙虚线：排除区间的滤波曲线。底部蓝带：吸气；黄带：呼气；空白：未分箱。竖／横虚线：分箱边界。','Caption')
    t=table(['呼吸状态','时间分箱','固定幅值分箱','核对数据占比'],widths)
    for row in rows[offset:offset+2]:
        cells=t.add_row().cells;cells[0].text=row['name']
        result=row['reference_comparison']
        for c,mode in zip(cells[1:3],['time','fixed']):
            p=c.paragraphs[0];p.paragraph_format.line_spacing=1;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(DATA/'figures'/f"{row['id']}_{mode}.png"),width=Inches(3.7 if offset==0 else 3.94))
            add_metric(c,result['modes'][mode]['mismatch_percent'])
        cells[3].text=fmt(result['scored_percent'])
        if result['scored_runs']<result['total_runs']:
            cells[3].add_paragraph(f"{result['scored_runs']}／{result['total_runs']} 次有核对数据")
        for p in cells[0].paragraphs+cells[3].paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    style_table(t,widths)
    if offset==0:
        para('各自参考规则不同，不据此作综合排名；核对占比低时仅代表少量片段。结果为模拟测算，不代表 CT 重建质量。','Caption')
    if any(row['id'] in ['apnea','periodic_weakening'] for row in rows[offset:offset+2]):
        para('模拟的暂停与周期性减弱仍保留微弱呼吸振荡，此处不评价暂停识别能力。','Caption')

d.core_properties.title='时间与固定幅值分箱波形对比'
d.core_properties.subject='模拟波形与分箱结果对照'
d.core_properties.author='RespiraScope'
d.save(OUT)
print(OUT)
